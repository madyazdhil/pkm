from __future__ import annotations

from collections import Counter
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "outputs/2026-10-03-dummy-datasets"
FILES = [
    "dummy-data-rekap-layanan-kecamatan.xlsx",
    "dummy-data-monitoring-surat-disposisi.xlsx",
    "dummy-data-agenda-kegiatan-internal.xlsx",
]


def read_table(path: Path):
    wb = load_workbook(path, data_only=True, read_only=True)
    assert wb.sheetnames == ["Data", "Panduan"], f"{path.name}: sheet order/name mismatch: {wb.sheetnames}"
    data = wb["Data"]
    rows = list(data.iter_rows(values_only=True))
    headers = [str(value or "").strip() for value in rows[0]]
    records = [dict(zip(headers, row)) for row in rows[1:] if any(value not in (None, "") for value in row)]
    guide = list(wb["Panduan"].iter_rows(min_row=6, max_row=15, max_col=3, values_only=True))
    guide_names = [row[0] for row in guide if row[0]]
    assert len(records) == 100, f"{path.name}: expected 100 records, got {len(records)}"
    assert guide_names == headers[:len(guide_names)], f"{path.name}: Panduan names do not preserve data headers"
    return headers, records


for file in FILES:
    headers, records = read_table(DATA / file)
    print(f"PASS structure {file}: {len(records)} records")

headers, records = read_table(DATA / FILES[0])
assert Counter(row["Status"] for row in records) == Counter({"Selesai": 25, "Diproses": 25, "Menunggu Dokumen": 25, "Dibatalkan": 25})
assert Counter(row["Jenis_Layanan"] for row in records) == Counter({
    "Administrasi Umum": 20,
    "Surat Keterangan": 20,
    "Rekomendasi": 20,
    "Informasi Publik": 20,
    "Pengaduan Non-Sensitif": 20,
})
assert all(row["Jumlah_Permohonan"] == 1 for row in records)
print("PASS primary distribution: KPI baseline 100 / 25 / 50 and five services × 20")

headers, records = read_table(DATA / FILES[1])
assert len({row["Prioritas"] for row in records}) == 3, "surat priorities should vary"
assert all((row["Tanggal_Selesai"] not in (None, "")) if row["Status_Disposisi"] == "Selesai" else (row["Tanggal_Selesai"] in (None, "")) for row in records), "surat completion dates inconsistent with status"
print("PASS surat semantics: varied priorities and completion dates only for completed records")

headers, records = read_table(DATA / FILES[2])
expected_pairs = {
    "Rapat Koordinasi": "Rapat Internal",
    "Sosialisasi Layanan": "Sosialisasi",
    "Evaluasi Bulanan": "Evaluasi",
    "Pelatihan Internal": "Pelatihan",
    "Kegiatan Administrasi": "Administrasi",
}
assert all(expected_pairs[row["Nama_Kegiatan"]] == row["Jenis_Kegiatan"] for row in records)
ref = date(2026, 10, 4)
for row in records:
    day = row["Tanggal"].date() if hasattr(row["Tanggal"], "date") else row["Tanggal"]
    status = row["Status_Persiapan"]
    if status == "Selesai": assert day < ref
    if status == "Berjalan": assert ref <= day <= date(2026, 10, 10)
    if status == "Siap": assert day > ref
    if status == "Belum Dimulai": assert day > ref
print("PASS agenda semantics: activity/type pairs and date/status coherence")
