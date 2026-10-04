from __future__ import annotations

import json
from collections import Counter
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "outputs/2026-10-03-dummy-datasets/dummy-data-rekap-layanan-kecamatan.xlsx"
OUT = ROOT / "materials/workshop/starter-dashboard/fixtures"


def serialise(value):
    if isinstance(value, (datetime, date)):
        return value.isoformat()[:10]
    return value


wb = load_workbook(SOURCE, data_only=True, read_only=True)
ws = wb["Data"]
rows = list(ws.iter_rows(values_only=True))
headers = [str(value or "").strip() for value in rows[0]]
records = [dict(zip(headers, [serialise(value) for value in row])) for row in rows[1:] if any(value not in (None, "") for value in row)]

fixture = {
    "source": SOURCE.name,
    "sheet": "Data",
    "records": records,
}
(OUT / "rekap-layanan.json").write_text(json.dumps(fixture, ensure_ascii=False, indent=2) + "\n")

status = Counter(record["Status"] for record in records)
service = Counter(record["Jenis_Layanan"] for record in records)
expected = {
    "source": SOURCE.name,
    "rowCount": len(records),
    "kpis": {
        "total": sum(int(record["Jumlah_Permohonan"]) for record in records),
        "completed": sum(int(record["Jumlah_Permohonan"]) for record in records if record["Status"] == "Selesai"),
        "inProcess": sum(int(record["Jumlah_Permohonan"]) for record in records if record["Status"] in {"Diproses", "Menunggu Dokumen"}),
    },
    "statusCounts": dict(sorted(status.items())),
    "serviceCounts": dict(sorted(service.items())),
}
(OUT / "expected-dashboard.json").write_text(json.dumps(expected, ensure_ascii=False, indent=2) + "\n")
print(json.dumps(expected, ensure_ascii=False))
